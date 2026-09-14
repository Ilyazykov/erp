"""
Damodaran's genuine monthly Implied ERP -- "ERP (T12m)" -- for direct visual
comparison against Shiller's Excess CAPE Yield on the US ERP tab.

NOTE (correcting an earlier version of this script): Damodaran's official
*annual* spreadsheet (histimpl.xls) is NOT the right source for a monthly
comparison -- an earlier version of this script stepped that annual figure
flat across months, which does not match Damodaran's own real monthly
series and reads as a staircase rather than the genuinely-varying line his
own data shows. Damodaran in fact recomputes and republishes ERP every
month (since September 2008), using trailing-twelve-month dividends +
buybacks against the current S&P 500 level and Treasury rate -- see his
blog post "The Price of Risk: An Equity Risk Premium Monologue!"
(aswathdamodaran.blogspot.com): "Rather than compute the implied equity
risk premium at the start of every year... I shifted to computing the
equity risk premium for the S&P 500 at the start of every month, in
September 2008." This script reads that genuine monthly series directly.

Source (no auth required): Damodaran's own NYU Stern page
    https://pages.stern.nyu.edu/~adamodar/pc/implprem/ERPbymonth.xlsx
    ("Historical ERP" sheet; "ERP (T12m)" column, one row per month,
    2008-09-present)

Z-score: rolling 36-month window (same pattern as us_erp.py / the rest of
the site), computed on this series' own history alone -- Damodaran's ERP
(T12m) only goes back to 2008-09, a much shorter and more recent window
than Shiller's (1871-present), so this Z-score has "seen" far fewer regimes
(no pre-2008 recessions, no dot-com crash) and should be read as a second,
independent, purely informational signal -- NOT used to drive the site's
US stocks/bonds allocation donut, which remains driven by Shiller's Z-score
only.

Outputs:
    data/damodaran_erp.csv   (date [YYYY-MM], implied_erp_t12m_pct, z_score)

Run:
    python erp_valuation/damodaran_erp.py
"""
from __future__ import annotations

import csv
import math
import urllib.request
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = REPO_ROOT / "data"
OUT_PATH = DATA_DIR / "damodaran_erp.csv"

DAMODARAN_URL = "https://pages.stern.nyu.edu/~adamodar/pc/implprem/ERPbymonth.xlsx"
SHEET_NAME = "Historical ERP"
DATE_COL = 1  # "Start of month"
ERP_T12M_COL = 10  # "ERP (T12m)" -- distinct from "ERP (Smoothed)"/"ERP (Normalized)"/etc.

ZSCORE_WINDOW = 36


def fetch_damodaran_xlsx() -> bytes:
    headers = {"User-Agent": "curl/8.0"}
    req = urllib.request.Request(DAMODARAN_URL, headers=headers)
    with urllib.request.urlopen(req, timeout=30) as r:
        return r.read()


def parse_monthly_erp(xlsx_bytes: bytes) -> list[dict]:
    """Returns list of {date (YYYY-MM), implied_erp_t12m_pct}, sorted by date."""
    tmp_path = DATA_DIR / "_damodaran_erp_tmp.xlsx"
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    tmp_path.write_bytes(xlsx_bytes)
    try:
        wb = openpyxl.load_workbook(tmp_path, data_only=True)
        sheet = wb[SHEET_NAME]

        rows = []
        for r in range(2, sheet.max_row + 1):
            date_val = sheet.cell(row=r, column=DATE_COL).value
            erp_val = sheet.cell(row=r, column=ERP_T12M_COL).value
            if date_val is None or erp_val is None:
                continue
            if not isinstance(erp_val, (int, float)):
                continue
            rows.append({
                "date": f"{date_val.year:04d}-{date_val.month:02d}",
                "implied_erp_t12m_pct": round(erp_val * 100, 4),
            })
        rows.sort(key=lambda x: x["date"])
        return rows
    finally:
        tmp_path.unlink(missing_ok=True)


def rolling_zscore(series: list[float | None], window: int = ZSCORE_WINDOW) -> list[float | None]:
    result: list[float | None] = []
    for i in range(len(series)):
        if series[i] is None:
            result.append(None)
            continue
        start = max(0, i - window)
        hist = [v for v in series[start:i] if v is not None]
        if len(hist) < 6:
            result.append(None)
            continue
        mean = sum(hist) / len(hist)
        variance = sum((v - mean) ** 2 for v in hist) / len(hist)
        std = math.sqrt(variance)
        result.append(0.0 if std < 1e-9 else (series[i] - mean) / std)
    return result


def main() -> int:
    print("Fetching Damodaran's ERPbymonth.xlsx (ERP T12m, monthly)...")
    xlsx_bytes = fetch_damodaran_xlsx()
    rows = parse_monthly_erp(xlsx_bytes)
    print(f"  Parsed {len(rows)} monthly rows ({rows[0]['date']} to {rows[-1]['date']})")

    series = [r["implied_erp_t12m_pct"] for r in rows]
    zscores = rolling_zscore(series)

    out_rows = [
        {
            "date": r["date"],
            "implied_erp_t12m_pct": r["implied_erp_t12m_pct"],
            "z_score": round(z, 4) if z is not None else "",
        }
        for r, z in zip(rows, zscores)
    ]

    with OUT_PATH.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["date", "implied_erp_t12m_pct", "z_score"])
        w.writeheader()
        w.writerows(out_rows)

    print(f"Saved {OUT_PATH}")
    populated = sum(1 for r in out_rows if r["z_score"] != "")
    print(f"Total rows: {len(out_rows)}, with z-score: {populated}")
    last = out_rows[-1]
    print(f"Latest ({last['date']}): ERP (T12m) = {last['implied_erp_t12m_pct']}%  Z = {last['z_score']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
